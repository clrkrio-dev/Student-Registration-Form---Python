import tkinter as tk
from tkinter import ttk, filedialog
from tkinter.messagebox import showwarning, showinfo, showerror
from PIL import Image, ImageTk
import openpyxl
import os

file_path = r"C:\Users\johnc\Documents\Student Registration Form.xlsx"
if not os.path.exists(file_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    headers = ["Student Number", "Student Name", "Father Name", "Mother Name", "Date of Birth", "Mobile Number",
               "Email Address", "Password", "Gender", "Course-Subject", "Photo", "Address"]
    sheet.append(headers)
    workbook.save(file_path)

workbook = openpyxl.load_workbook(file_path)
sheet = workbook.active

ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "12345"

root = tk.Tk()
root.title("Student Registration Form")
root.geometry("400x400")
root.resizable(False, False)

login_frame = tk.Frame(root)
login_frame.pack(pady=50)

tk.Label(login_frame, text="ADMIN LOGIN", font=("Times New Roman", 18)).pack(pady=10)
tk.Label(login_frame, text="Username", font=("Times New Roman", 12)).pack()
username_entry = tk.Entry(login_frame, font=("Times New Roman", 12), width=15)
username_entry.pack()
tk.Label(login_frame, text="Password:", font=("Times New Roman", 12)).pack()
password_frame = tk.Frame(login_frame)
password_frame.pack()
password_entry = tk.Entry(password_frame, font=("Times New Roman", 12), show="*", width=10)
password_entry.pack(side="left")
def toggle_password():
    if password_entry.cget('show') == "":
        password_entry.config(show="*")
        show_btn.config(text="Show")
    else:
        password_entry.config(show="")
        show_btn.config(text="Hide")
show_btn = tk.Button(password_frame, text="Show", command=toggle_password)
show_btn.pack(side="left", padx=5)

def open_dashboard():
    login_frame.pack_forget()
    root.geometry("1150x700")
    root.title("Student Registration Form")
    root.resizable(True, True)
    main_frame.pack(fill="both", expand=1)

def login():
    user = username_entry.get()
    pwd = password_entry.get()
    if user == "" or pwd == "":
        showwarning("Missing", "Please fill all fields")
        return
    if user == ADMIN_USERNAME and pwd == ADMIN_PASSWORD:
        showinfo("Success", "Login Successful!")
        open_dashboard()
    else:
        showerror("Error", "Invalid Username or Password")

tk.Button(login_frame, text="Login", font=("Arial", 12), width=10, command=login).pack(pady=15)

main_frame = tk.Frame(root)
main_frame.pack_forget()

left_frame = tk.Frame(main_frame)
left_frame.pack(side="left", fill="both", expand=1)
right_frame = tk.Frame(main_frame, width=420)
right_frame.pack(side="right", fill="y")
canvas = tk.Canvas(left_frame)
canvas.pack(side="left", fill="both", expand=1)
vertical_scrollbar = ttk.Scrollbar(left_frame, orient="vertical", command=canvas.yview)
vertical_scrollbar.pack(side="right", fill="y")
canvas.configure(yscrollcommand=vertical_scrollbar.set)
scrollbar_frame = tk.Frame(canvas)
canvas.create_window((0, 0), window=scrollbar_frame, anchor="nw")
def on_configure(e):
    canvas.configure(scrollregion=canvas.bbox("all"))
scrollbar_frame.bind("<Configure>", on_configure)

Title = tk.Frame(scrollbar_frame)
Title.pack()
tk.Label(Title, text="Student Registration Form", font=("Times New Roman", 15, "bold")).pack(side="top", pady=5, padx=10)

def add_placeholder(entry, placeholder):
    entry.insert(0, placeholder)
    entry.config(fg="gray")
    def on_focus_in(event):
        if entry.get() == placeholder:
            entry.delete(0, tk.END)
            entry.config(fg="black")
    def on_focus_out(event):
        if entry.get().strip() == "":
            entry.delete(0, tk.END)
            entry.insert(0, placeholder)
            entry.config(fg="gray")
    entry.bind("<FocusIn>", on_focus_in)
    entry.bind("<FocusOut>", on_focus_out)

def validate_input(new_value):
    return new_value.isdigit() or new_value == ""
vcmd = (scrollbar_frame.register(validate_input), "%P")

Number_frame = tk.Frame(scrollbar_frame)
Number_frame.pack(pady=5, fill="x", padx=10)
tk.Label(Number_frame, text="Student Number: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Number_entry = tk.Entry(Number_frame, validate="key", width=40, validatecommand=vcmd)
Number_entry.pack(side="left", padx=5)

Name_frame = tk.Frame(scrollbar_frame)
Name_frame.pack(pady=5, fill="x", padx=10)
tk.Label(Name_frame, text="Student Name: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Name_entry_1 = tk.Entry(Name_frame, width=18)
add_placeholder(Name_entry_1, "First Name")
Name_entry_1.pack(side="left", padx=5)
Name_entry_2 = tk.Entry(Name_frame, width=18)
add_placeholder(Name_entry_2, "Last Name")
Name_entry_2.pack(side="left", padx=5)

Father_Name_Frame = tk.Frame(scrollbar_frame)
Father_Name_Frame.pack(pady=5, fill="x", padx=10)
tk.Label(Father_Name_Frame, text="Father Name: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Father_entry = tk.Entry(Father_Name_Frame, width=50)
Father_entry.pack(side="left", padx=5)

Mother_Name_Frame = tk.Frame(scrollbar_frame)
Mother_Name_Frame.pack(pady=5, fill="x", padx=10)
tk.Label(Mother_Name_Frame, text="Mother Name:", font=("Times New Roman", 12)).pack(side="left", padx=5)
Mother_entry = tk.Entry(Mother_Name_Frame, width=50)
Mother_entry.pack(side="left", padx=5)

DOB_Frame = tk.Frame(scrollbar_frame)
DOB_Frame.pack(pady=5, fill="x", padx=10)
tk.Label(DOB_Frame, text="Date of Birth: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Day_entry = tk.Entry(DOB_Frame, width=6)
add_placeholder(Day_entry, "Day")
Day_entry.pack(side="left", padx=5)
Month_entry = tk.Entry(DOB_Frame, width=8)
add_placeholder(Month_entry, "Month")
Month_entry.pack(side="left", padx=5)
Year_entry = tk.Entry(DOB_Frame, width=10)
add_placeholder(Year_entry, "Year")
Year_entry.pack(side="left", padx=5)
tk.Label(DOB_Frame, text="(DD/MM/YY)", font=("Times New Roman", 10, "italic")).pack(side="left", padx=5)

vcmd_1 = (scrollbar_frame.register(validate_input), "%P")

Phone_Frame = tk.Frame(scrollbar_frame)
Phone_Frame.pack(padx=10, fill="x", pady=5)
tk.Label(Phone_Frame, text="Contant Number: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Code_entry = tk.Entry(Phone_Frame, width=6)
add_placeholder(Code_entry, "+63")
Code_entry.pack(side="left", padx=5)
Phone_entry = tk.Entry(Phone_Frame, width=30, validate="key", validatecommand=vcmd_1)
Phone_entry.pack(side="left", padx=5)

Email_Frame = tk.Frame(scrollbar_frame)
Email_Frame.pack(padx=10, fill="x", pady=5)
tk.Label(Email_Frame, text="Email Address: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Email_entry = tk.Entry(Email_Frame, width=50)
Email_entry.pack(side="left", padx=5)
add_placeholder(Email_entry, "Helloworld123@gmail.com")

Password_Frame = tk.Frame(scrollbar_frame)
Password_Frame.pack(padx=10, pady=5, fill="x")
tk.Label(Password_Frame, text="Password: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Password_entry = tk.Entry(Password_Frame, width=40, show="*")
Password_entry.pack(side="left", padx=5)
def Button_password():
    if Password_entry.cget("show") == "*":
        Password_entry.config(show="")
        Password_Button.config(text="Hide")
    else:
        Password_entry.config(show="*")
        Password_Button.config(text="Show")
Password_Button = tk.Button(Password_Frame, text="Show", command=Button_password)
Password_Button.pack(side="left", padx=5)

Gender_Frame = tk.Frame(scrollbar_frame)
Gender_Frame.pack(padx=10, pady=5, fill="x")
Selected_option = tk.StringVar()
Selected_option.set("Male")
tk.Label(Gender_Frame, text="Gender: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
tk.Radiobutton(Gender_Frame, text="Male", value="Male", variable=Selected_option).pack(side="left", padx=10)
tk.Radiobutton(Gender_Frame, text="Female", value="Female", variable=Selected_option).pack(side="left", padx=10)

Course_Frame = tk.Frame(scrollbar_frame)
Course_Frame.pack(padx=10, fill="x", pady=5)
tk.Label(Course_Frame, text="Course: ", font=("Times New Roman", 12)).pack(side="left", padx=5)

IT_var = tk.IntVar()
CS_var = tk.IntVar()
COE_var = tk.IntVar()
ELE_var = tk.IntVar()
all_var = (IT_var, CS_var, COE_var, ELE_var)
all_course = ['IT', 'CS', 'COE', 'ELE']

Course_Subject = { 
    "IT": { "1st Year":['Fundemental Concept of Mathmatics', 'English Plus', 'Computer Programming 1','NSTP 1','Purposive Communication','Art Appreciation', 'Science, Technology, and Society', 'Ethics', 'General Elective', 'Introduction to Computing', 'Self Testing Activities', 'Computer Programming 2','NSTP 2', 'Discrete Mathematics', 'Understanding the Self', 'Readings in Philippines History', 'The Contemporary World', 'Intro to Human & Computer Interaction', 'Rhythmic Activities'], 
            "2nd Year":['Data Structures & Algorithms', 'General Elective 2 (The Entrepreneurial Mind)', 'General Elective 3 (Philippine Popular Culture)', 'Elective 1 (Object Oriented Programming)', 'Games and Sport', 'Elective 2 (Platform Technologies)', 'Mathematics in the Modern World', 'Rizal Life and Works', 'Information Management 1', 'Integrative Programming and Technologies', 'Networking 1', 'Recreational Activities', 'Quantitative Methods (Incl. Modeling & Simulation)'], 
            "3rd Year":['Advanced Database System', 'Networking 2', 'System Integration and Architecture 1', 'Elective 3 (Wed System and Technologies)', 'Applications Development and Emerging Technologies','Information Assurance and Security 1', 'System Administrations adn Maintenance', 'Social and Professional Issues'],
            "4th Year":['Information Assurance and Security 2', 'Capstone Project 1', 'Practicum', 'Elective 4 (System Integration & Architecture 2)', ' Capstone Project 2'], }, 
    "CS": { "1st Year":['Fundemental Concept of Mathmatics', 'English Plus', 'Computer Programming 1','NSTP 1','Purposive Communication','Art Appreciation', 'Science, Technology, and Society', 'Ethics', 'General Elective', 'Introduction to Computing', 'Self Testing Activities', 'NSTP2','Computer Programming 2 (Intermediate Programming)', 'Discrete Structure 1', 'Understanding the Self', 'Readings in Philippine History', 'The Contemporary World', 'Intro to Human & Computer Interaction', 'Rhythmic Activities'], 
           "2nd Year":['Algorithms & Complexity', ' Discrete Structures 2', 'Data Structure & Algorithms', 'General Elective 2 (The Entrepreneurial Mind)', 'General Elective 3 (Philippine Popular Culture)', 'Object Oriented Programming', 'Games and Sport', 'Mathematics in the Modern World', 'Rizal Life and Works', 'Information Management 1', 'Math Elective (Calculus 1)', ' Recreational Activities'], 
           "3rd Year":['Application Development and Emerging Technologies', 'Architecture and Organization', ' CS Elective 1 (Graphics & Visual Computin', 'Automation Theory And Formal Languages', 'Information Assurance And Security', 'CS Elective 2 (Parallel & Distributed Computing)', 'CS Elective( Intelligent Systems)', 'Programming Languages', ' Opearting Systems', 'Social Issues and Professional Practice 1', 'Software Enginerring 1'],
            "4th Year":['Practicum', 'Thesis 1', ' Software Engineering 2', 'Thesis 2', 'Network and Communications'], }, 
    "COE": { "1st Year":['Fundemental Concept of Mathmatics', 'English Plus', 'Calculus 1', 'Computer Engineering as a Discipline', 'Chemistry for Engineers', 'Understanding the Self', 'Science, Technology, and Society', 'Mathematics in the Modern World', 'Programming Logic and Design', 'Self Testing Activities', 'NSTP1', 'Calculus 2', ' Physiscs for Engineers', 'Object Oriented Programming', 'Readings in Philippine History', 'Enginnering Data Analysis', 'Discrete Mathematics', 'Rhthmis Activities', 'NSTP2'], 
            "2nd Year":['Differential Equations', 'Art Appreciation', 'Data Structure and Algorithms', 'Engineering Economics', 'Fundemental of Electrical Circuits', 'General Elective (Living in the IT Era)', 'Computer Aided Drafting', 'Games and Sport', 'Numerical Methods', 'Software Design', 'Purposive Communication', 'Fundemental of Electronics Circuits', 'Rizal Life and Works', 'The Contemporary World', ' Recreational Activities'], 
            "3rd Year":['Logic Circuits and Designs', 'Operating Systems', 'Data and Digital Communications', 'Introduction to HDL', 'Feedback and Control System', 'Fundemental of Mixed Signals and Sensors', ' Computer Engineering Drafting and Design', 'Cognate/Elective Course 1', 'Basic Occupitional Health and Safety', ' Computer Networks and Security', 'Microprocessor', 'Methods of Research', 'Technoprenership', 'Ethics', 'CpE Laws and Professional Practice', 'Cognate/Elecive Course 2'],
            "4th Year":['Embedded Systems', 'Computer Architecture and Organization', 'Emerging Technologies in CpE', 'Cpe Practice and Design 1', 'Digital Signal Processing', 'General Elecctive 2 ( The Entrepreneural Mind)', 'Cognate/Elective Course 3', 'CpE Practice and Design 2', 'Seminars and Fieldtrips', ' On the Job Training(240 hrs)', ' General Elective 3 (Philippines Popular Culture)'], },
    "ELE": { "1st Year":['Fundemental Concept of Mathmatics', 'English Plus', 'Calculus 1','NSTP 1', 'Engineering Data Analysis','Chemistry for Engineers','Understanding the Self','Mathematics for the Modern World', 'Science, Technology, and Society', 'Ethics', 'Self Testing Activities', 'Calculus 2', 'Computer Programming', 'NSTP2', 'Physics for Engineers 1', 'Physics for Engineers 2', 'Readings in Philippine History', 'Material Science and Engineering', 'Rhythimic Activities'], 
            "2nd Year":['Computer Aided Drafting', 'Circuits 1', 'Differential Equations', 'ECE Laws, Contracts, Ethics, Standard & Safety', 'Art Appreciation', 'General Elective (Living in the IT Era)', 'Games and Sports', 'Advanced Engineering Mathematics For ECE', 'Circuits 2', 'Electronics 2: Electronis Circuit Analysis and Design', 'Electromagnetics', 'Rizal Life and Works', 'Recreational Activities', 'Communication 1 : Principles of Communications System'], 
            "3rd Year":['Electronics 3: Electronics Systems and Design', 'Contemporary World', 'Purposive Communication', 'Ethics', 'Digital Electronics 1: Logic Circuit & Switching Theory', 'Communication 2: Modulation and Coding Techniques',' Signals, Spectra & Signal Processing', 'Communication 3: Transmission Media & Antenna System & Design', 'Communication 4: Data Communication', 'Feedback and Control System', 'General Elective 2 (The Entrepreneural Mind)', 'Digital Electronics 2: Microprocessor & Microcontroller System & Design', 'Technopreneuship'], 
            "4th Year":['Design 1/Capstone Project 1', 'ECE Elective 1: Advanced Communication System & Design(Wireless)', 'Environmental Science and Engineering', 'ECE Elective: Robotics Technology', 'General Elective 3 (Philippine Popular Culture)', 'Methods of Research', 'Comprehensive Class','Design 2/Capstone Project 2', 'ECE Elective 2: Advanced Networking', 'On the Job Training (240 hrs)', 'Seminars and Fieldtrips'],
              }
         }

def check_selected(selected_var):
    for var in all_var:
        if var != selected_var:
            var.set(0)
    update_combobox()

def update_combobox():
    Box.set("-------------Year Level-------------")
    Listbox.delete(0, tk.END)
    selected_course = ""
    for var, course_name in zip(all_var, all_course):
        if var.get() == 1:
            selected_course = course_name
            break
    if selected_course == "":
        Box["values"] = []
        return
    Box["values"] = [f"{selected_course}-1st Year", f"{selected_course}-2nd Year",
                     f"{selected_course}-3rd Year", f"{selected_course}-4th Year"]

for var, course_name in zip(all_var, all_course):
    tk.Checkbutton(Course_Frame, text=course_name, variable=var, command=lambda v=var: check_selected(v)).pack(side="left", padx=8)

Subject_frame = tk.Frame(scrollbar_frame)
Subject_frame.pack(fill="x", padx=10, pady=5)
tk.Label(Subject_frame, text="Year Level:", font=("Times New Roman", 12)).pack(side="left", padx=5)
Box = ttk.Combobox(Subject_frame, width=30, state="readonly")
Box.set("-------------Year Level-------------")
Box.pack(side="left", padx=5)
Listbox_Frame = tk.Frame(scrollbar_frame)
Listbox_Frame.pack(fill="x", pady=5, padx=10)
tk.Label(Listbox_Frame, text="List of Subject: ", font=("Times New Roman", 12)).pack(side="left", padx=5)
Listbox = tk.Listbox(Listbox_Frame, height=8, width=50, selectmode="multiple")
Listbox.pack(side="left", padx=5)

def show_subjects(event):
    Listbox.delete(0, tk.END)
    val = Box.get()
    if val == "-------------Year Level-------------":
        return
    course, year = val.split("-")
    subjects = Course_Subject[course][year]
    for sub in subjects:
        Listbox.insert(tk.END, sub)
Box.bind("<<ComboboxSelected>>", show_subjects)

Photo_var = tk.StringVar()
def open_file():
    filepath = filedialog.askopenfilename(title="Select Student Photo", filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.gif")])
    if filepath:
        Photo_var.set(filepath)
        File_label.config(text=os.path.basename(filepath))
        img = Image.open(filepath)
        img = img.resize((100, 100))
        Photo_img = ImageTk.PhotoImage(img)
        Photo_label_preview.image = Photo_img
        Photo_label_preview.config(image=Photo_img)
    else:
        Photo_var.set("")
        File_label.config(text="No file Chosen")
        Photo_label_preview.config(image="")

Studentphoto_frame = tk.Frame(scrollbar_frame)
Studentphoto_frame.pack(fill="x", padx=10, pady=5)
tk.Label(Studentphoto_frame, text="Student Photo:", font=("Times New Roman", 12)).pack(side="left", padx=5)
Photo_label_preview = tk.Label(Studentphoto_frame)
Photo_label_preview.pack(side="left", padx=5)
tk.Button(Studentphoto_frame, text="Choose File", command=open_file).pack(side="left", padx=5)
File_label = tk.Label(Studentphoto_frame, text="No file Chosen")
File_label.pack(side="left", padx=5)

Studentaddress_frame = tk.Frame(scrollbar_frame)
Studentaddress_frame.pack(fill="x", padx=10, pady=5)
tk.Label(Studentaddress_frame, text="Address", font=("Times New Roman", 12)).pack(side="left", padx=5)
Studentaddress_entry = tk.Entry(Studentaddress_frame, width=40)
Studentaddress_entry.pack(side="left", padx=5)

Filter_Frame = tk.Frame(right_frame)
Filter_Frame.pack(fill="x", padx=5, pady=5)
tk.Label(Filter_Frame, text="Search:", font=("Times New Roman", 11)).pack(side="left", padx=5)
Search_entry = tk.Entry(Filter_Frame, width=25)
Search_entry.pack(side="left", padx=5)

Tree_frame = tk.Frame(right_frame)
Tree_frame.pack(fill="both", expand=1, padx=5, pady=5)
Student_Tree = ttk.Treeview(Tree_frame, columns=("Student Number","Student Name","Father Name","Mother Name","Date of Birth","Mobile Number","Email Address","Password","Gender","Course-Subject","Address"), show="headings")
for col in Student_Tree["columns"]:
    Student_Tree.heading(col, text=col)
    Student_Tree.column(col, width=120, anchor="center")
Student_Tree.pack(fill="both", expand=1)

def load_treeview(filter_text=""):
    for i in Student_Tree.get_children():
        Student_Tree.delete(i)
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    filter_text = filter_text.lower()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = (row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8], row[9], row[11])
        if filter_text:
            if filter_text not in str(row[0]).lower() and filter_text not in str(row[1]).lower() and filter_text not in str(row[9]).lower():
                continue
        Student_Tree.insert("", tk.END, values=row_data)

def search_tree(event):
    load_treeview(Search_entry.get())

Search_entry.bind("<KeyRelease>", search_tree)

def clear_form():
    Number_entry.delete(0, tk.END)
    Name_entry_1.delete(0, tk.END)
    Name_entry_2.delete(0, tk.END)
    Father_entry.delete(0, tk.END)
    Mother_entry.delete(0, tk.END)
    Day_entry.delete(0, tk.END)
    Month_entry.delete(0, tk.END)
    Year_entry.delete(0, tk.END)
    Code_entry.delete(0, tk.END)
    Phone_entry.delete(0, tk.END)
    Email_entry.delete(0, tk.END)
    Password_entry.delete(0, tk.END)
    for var in all_var:
        var.set(0)
    Selected_option.set("Male")
    Box.set("-------------Year Level-------------")
    Listbox.delete(0, tk.END)
    File_label.config(text="No file Chosen")
    Photo_label_preview.config(image="")
    Studentaddress_entry.delete(0, tk.END)

def submit_student():
    number = Number_entry.get()
    name = Name_entry_1.get() + " " + Name_entry_2.get()
    father = Father_entry.get()
    mother = Mother_entry.get()
    dob = f"{Day_entry.get()}/{Month_entry.get()}/{Year_entry.get()}"
    mobile = Code_entry.get() + Phone_entry.get()
    email = Email_entry.get()
    password = Password_entry.get()
    gender = Selected_option.get()
    course = Box.get()
    if course != "-------------Year Level-------------":
        selected_subjects = ", ".join([Listbox.get(i) for i in Listbox.curselection()])
        course_subject = f"{course}: {selected_subjects}"
    else:
        course_subject = ""
    photo = Photo_var.get()
    address = Studentaddress_entry.get()
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    sheet.append([number, name, father, mother, dob, mobile, email, password, gender, course_subject, photo, address])
    workbook.save(file_path)
    load_treeview()
    showinfo("Success", "Student Added Successfully!")
    clear_form()

def update_student():
    selected_item = Student_Tree.selection()
    if not selected_item:
        showwarning("Warning", "Please select a student to update")
        return
    item = Student_Tree.item(selected_item)
    values = item['values']
    number = Number_entry.get()
    name = Name_entry_1.get() + " " + Name_entry_2.get()
    father = Father_entry.get()
    mother = Mother_entry.get()
    dob = f"{Day_entry.get()}/{Month_entry.get()}/{Year_entry.get()}"
    mobile = Code_entry.get() + Phone_entry.get()
    email = Email_entry.get()
    password = Password_entry.get()
    gender = Selected_option.get()
    course = Box.get()
    if course != "-------------Year Level-------------":
        selected_subjects = ", ".join([Listbox.get(i) for i in Listbox.curselection()])
        course_subject = f"{course}: {selected_subjects}"
    else:
        course_subject = ""
    photo = Photo_var.get()
    address = Studentaddress_entry.get()
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == str(values[0]):
            row[0].value = number
            row[1].value = name
            row[2].value = father
            row[3].value = mother
            row[4].value = dob
            row[5].value = mobile
            row[6].value = email
            row[7].value = password
            row[8].value = gender
            row[9].value = course_subject
            row[10].value = photo
            row[11].value = address
            break
    workbook.save(file_path)
    load_treeview()
    showinfo("Success", "Student Updated Successfully!")
    clear_form()

def on_tree_select(event):
    selected_item = Student_Tree.selection()

    if not selected_item:
        return
    item = Student_Tree.item(selected_item)
    values = item['values']
    Number_entry.delete(0, tk.END)
    Number_entry.insert(0, values[0])
    full_name = values[1].split(" ")
    Name_entry_1.delete(0, tk.END)
    Name_entry_1.insert(0, full_name[0])
    Name_entry_2.delete(0, tk.END)

    if len(full_name) > 1:
        Name_entry_2.insert(0, " ".join(full_name[1:]))
    Father_entry.delete(0, tk.END)
    Father_entry.insert(0, values[2])
    Mother_entry.delete(0, tk.END)
    Mother_entry.insert(0, values[3])
    day, month, year = values[4].split("/")
    Day_entry.delete(0, tk.END)
    Day_entry.insert(0, day)
    Month_entry.delete(0, tk.END)
    Month_entry.insert(0, month)
    Year_entry.delete(0, tk.END)
    Year_entry.insert(0, year)
    Code_entry.delete(0, tk.END)
    Phone_entry.delete(0, tk.END)

    if values[5].startswith("+63"):
        Code_entry.insert(0, "+63")
        Phone_entry.insert(0, values[5][3:])
    else:
        Phone_entry.insert(0, values[5])
    Email_entry.delete(0, tk.END)
    Email_entry.insert(0, values[6])
    Password_entry.delete(0, tk.END)
    Password_entry.insert(0, values[7])
    Selected_option.set(values[8])
    
    for var, course_name in zip(all_var, all_course):
        if values[9].startswith(course_name):
            var.set(1)
        else:
            var.set(0)
    update_combobox()
    if ":" in values[9]:
        course_part, subjects_part = values[9].split(":")
        Box.set(course_part.strip())
        show_subjects(None)
        Listbox.selection_clear(0, tk.END)
        subjects = [s.strip() for s in subjects_part.split(",")]
        for i in range(Listbox.size()):
            if Listbox.get(i) in subjects:
                Listbox.selection_set(i)
    else:
        Box.set("-------------Year Level-------------")
        Listbox.delete(0, tk.END)
    Photo_var.set(values[10])
    Studentaddress_entry.delete(0, tk.END)
    Studentaddress_entry.insert(0, values[11])
    if Photo_var.get() and os.path.exists(Photo_var.get()):
        img = Image.open(Photo_var.get())
        img = img.resize((100,100))
        Photo_img = ImageTk.PhotoImage(img)
        Photo_label_preview.image = Photo_img
        Photo_label_preview.config(image=Photo_img)
    else:
        Photo_label_preview.config(image="")

Student_Tree.bind("<<TreeviewSelect>>", on_tree_select)

Button_frame = tk.Frame(scrollbar_frame)
Button_frame.pack(pady=5)
tk.Button(Button_frame, text="Submit", command=submit_student, width=15).pack(side="left", padx=5)
tk.Button(Button_frame, text="Update", command=update_student, width=15).pack(side="left", padx=5)
tk.Button(Button_frame, text="Clear", command=clear_form, width=15).pack(side="left", padx=5)

load_treeview()
root.mainloop()
